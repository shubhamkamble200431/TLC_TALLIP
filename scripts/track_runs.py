#!/usr/bin/env python
"""
Live tracker for the 2^4 factorial runs (NLLB `factorial` + `mbart50`).

Scans experiments/v2/<exp>/<X-ABCD>/seed_<S>/ for status.json / metrics.json /
evaluation/v2/unsampled/metrics.json and writes:

    results/run_tracker.xlsx     (sheets: Runs, Summary)
    results/run_tracker.csv      (the Runs sheet, plain)

Covers ALL 96 planned runs (2 experiments x 16 configs x 3 seeds) — the 6 that
were already finished, the ones in flight, and the ones still queued — and marks
which SLURM job (run_a/run_b/run_c/run_d) owns each.  Re-run any time; it is
idempotent.  A crontab entry refreshes it every few minutes.

    python scripts/track_runs.py
"""
import _paths  # noqa: F401

import csv
import datetime as dt
import json
import os
import statistics as stats
import subprocess

from src.config import REPO_ROOT

EXPERIMENTS = ["factorial", "mbart50"]
SEEDS = [42, 43, 44]
N_EPOCHS = 20

# cost-balanced halves (2 configs per RoPE x ALiBi quadrant each)
HALF_A = ["X-1000", "X-0001", "X-0100", "X-0010", "X-1010", "X-1001", "X-0011", "X-1111"]
HALF_B = ["X-0000", "X-1100", "X-0110", "X-1110", "X-0101", "X-1101", "X-0111", "X-1011"]
JOB_OF = {  # (experiment, half) -> slurm job name
    ("factorial", "A"): "run_a", ("factorial", "B"): "run_b",
    ("mbart50", "A"): "run_c",   ("mbart50", "B"): "run_d",
}

TZ = dt.datetime.now().astimezone().tzinfo


def quadrant(cfg: str) -> str:
    c, d = cfg[4], cfg[5]          # X-ABCD  -> C=RoPE (idx4), D=ALiBi (idx5)
    return {("0", "0"): "plain", ("1", "0"): "RoPE",
            ("0", "1"): "ALiBi", ("1", "1"): "RoPE+ALiBi"}[(c, d)]


def _load(path):
    try:
        with open(path) as f:
            return json.load(f)
    except Exception:
        return None


def _fmt(iso):
    if not iso:
        return ""
    try:
        return dt.datetime.fromisoformat(iso).astimezone(TZ).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return iso


def _hours(a, b):
    try:
        return round((dt.datetime.fromisoformat(b) - dt.datetime.fromisoformat(a)).total_seconds() / 3600, 2)
    except Exception:
        return ""


def squeue_states():
    """{job_name: state} for this user's run_* jobs."""
    out = {}
    try:
        r = subprocess.run(["squeue", "-u", os.environ.get("USER", "kahuja"),
                            "-h", "-o", "%j|%T|%N"], capture_output=True, text=True, timeout=15)
        for line in r.stdout.splitlines():
            name, state, node = (line.split("|") + ["", "", ""])[:3]
            out[name] = state
    except Exception:
        pass
    return out


def collect():
    jobs = squeue_states()
    now = dt.datetime.now(TZ).isoformat()
    rows = []
    for exp in EXPERIMENTS:
        for half, configs in (("A", HALF_A), ("B", HALF_B)):
            job = JOB_OF[(exp, half)]
            jstate = jobs.get(job, "—")
            for ci, cfg in enumerate(configs):
                for si, seed in enumerate(SEEDS):
                    d = os.path.join(REPO_ROOT, "experiments", "v2", exp, cfg, f"seed_{seed}")
                    st = _load(os.path.join(d, "status.json")) or {}
                    mtr = _load(os.path.join(d, "metrics.json")) or {}
                    tst = _load(os.path.join(d, "evaluation", "v2", "unsampled", "metrics.json")) or {}
                    cfy = {}
                    try:
                        import yaml
                        with open(os.path.join(d, "config.yaml")) as f:
                            cfy = yaml.safe_load(f) or {}
                    except Exception:
                        pass

                    status = st.get("status", "not started")
                    started, ended = st.get("started_at"), st.get("ended_at")
                    if status == "running" and started and not ended:
                        dur = _hours(started, now)
                    else:
                        dur = _hours(started, ended) if (started and ended) else ""

                    # progress (last eval epoch / total)
                    prog = ""
                    lh = mtr.get("log_history", [])
                    evs = [x for x in lh if "eval_bleu" in x]
                    last_step = lh[-1].get("step") if lh else None
                    if evs:
                        prog = f"epoch {evs[-1].get('epoch', '?')}/{N_EPOCHS}"
                    elif status == "running" and last_step:
                        prog = f"step {last_step}"
                    elif status == "completed":
                        prog = f"epoch {N_EPOCHS}/{N_EPOCHS}"

                    val_bleu = round(evs[-1]["eval_bleu"], 2) if evs else ""
                    tm = tst.get("metrics", {})
                    pdb = cfy.get("per_device_train_batch_size", "")
                    ga = cfy.get("gradient_accumulation_steps", "")
                    eff = pdb * ga if (isinstance(pdb, int) and isinstance(ga, int)) else ""

                    rows.append({
                        "Experiment": exp,
                        "Job": job,
                        "JobState": jstate,
                        "Half": half,
                        "Order": ci * len(SEEDS) + si + 1,
                        "Config": cfg,
                        "Quadrant": quadrant(cfg),
                        "Seed": seed,
                        "Status": status,
                        "Start": _fmt(started),
                        "End": _fmt(ended),
                        "Duration_h": dur,
                        "Progress": prog,
                        "Val_BLEU": val_bleu,
                        "Test_BLEU": round(tm["bleu"], 2) if "bleu" in tm else "",
                        "Test_chrF2": round(tm["chrf2"], 2) if "chrf2" in tm else "",
                        "Test_COMET": round(tm["comet"], 2) if "comet" in tm else "",
                        "Eff_Batch": eff,
                        "Dir": os.path.relpath(d, REPO_ROOT),
                    })
    return rows


COLUMNS = ["Experiment", "Job", "JobState", "Half", "Order", "Config", "Quadrant", "Seed",
           "Status", "Start", "End", "Duration_h", "Progress", "Val_BLEU",
           "Test_BLEU", "Test_chrF2", "Test_COMET", "Eff_Batch", "Dir"]


def write_csv(rows, path):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Runs"

    fills = {
        "completed":   PatternFill("solid", fgColor="C6EFCE"),
        "running":     PatternFill("solid", fgColor="FFEB9C"),
        "failed":      PatternFill("solid", fgColor="FFC7CE"),
        "interrupted": PatternFill("solid", fgColor="F8CBAD"),
        "not started": PatternFill("solid", fgColor="F2F2F2"),
        "skipped":     PatternFill("solid", fgColor="D9D9D9"),
    }
    hdr = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="404040")

    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        ws.cell(row=1, column=c).font = hdr
        ws.cell(row=1, column=c).fill = hfill
        ws.cell(row=1, column=c).alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([r[c] for c in COLUMNS])
        fill = fills.get(r["Status"])
        if fill:
            ws.cell(row=ws.max_row, column=COLUMNS.index("Status") + 1).fill = fill

    ws.freeze_panes = "A2"
    widths = {"Dir": 42, "Start": 20, "End": 20, "Progress": 14, "Quadrant": 12,
              "Experiment": 11, "Status": 12, "JobState": 10}
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 11)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    # ---- Summary sheet -------------------------------------------------
    sm = wb.create_sheet("Summary")
    sm.append(["Generated", dt.datetime.now(TZ).strftime("%Y-%m-%d %H:%M:%S %Z")])
    sm.append([])
    sm.append(["Experiment", "Job", "Total", "Completed", "Running", "Not started",
               "Failed/Intr", "First start", "Last end", "Compute h"])
    for c in range(1, 11):
        sm.cell(row=3, column=c).font = Font(bold=True)

    def _grp(pred):
        g = [x for x in rows if pred(x)]
        done = [x for x in g if x["Status"] == "completed"]
        run = [x for x in g if x["Status"] == "running"]
        ns = [x for x in g if x["Status"] in ("not started", "skipped")]
        bad = [x for x in g if x["Status"] in ("failed", "interrupted")]
        starts = [x["Start"] for x in g if x["Start"]]
        ends = [x["End"] for x in done if x["End"]]
        comp = sum(x["Duration_h"] for x in done if isinstance(x["Duration_h"], (int, float)))
        return [len(g), len(done), len(run), len(ns), len(bad),
                min(starts) if starts else "", max(ends) if ends else "", round(comp, 1)]

    for exp in EXPERIMENTS:
        for job in sorted({r["Job"] for r in rows if r["Experiment"] == exp}):
            sm.append([exp, job] + _grp(lambda x, e=exp, j=job: x["Experiment"] == e and x["Job"] == j))
    sm.append(["ALL", "", *_grp(lambda x: True)])
    for c in range(1, 11):
        sm.cell(row=sm.max_row, column=c).font = Font(bold=True)
    for i, w in enumerate([12, 8, 8, 11, 9, 12, 12, 20, 20, 10], 1):
        sm.column_dimensions[get_column_letter(i)].width = w

    # ---- Aggregate sheet — mean/var over completed seeds, per config ----
    AGG_METRICS = [("Test_BLEU", "BLEU"), ("Test_chrF2", "chrF2"),
                   ("Test_COMET", "COMET"), ("Val_BLEU", "ValBLEU")]
    ag = wb.create_sheet("Aggregate")
    ag.append([f"Generated {dt.datetime.now(TZ).strftime('%Y-%m-%d %H:%M:%S %Z')}  "
                "— mean / sample variance over the seeds completed SO FAR (n<3 = partial)"])
    ag.append([])
    hdr = ["Experiment", "Config", "Quadrant", "n_seeds(/3)", "Eff_Batch"]
    for _, label in AGG_METRICS:
        hdr += [f"{label}_mean", f"{label}_var", f"{label}_sd"]
    ag.append(hdr)
    agg_hdr_font, agg_hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="404040")
    for c in range(1, len(hdr) + 1):
        ag.cell(row=3, column=c).font = agg_hdr_font
        ag.cell(row=3, column=c).fill = agg_hdr_fill

    def _agg_group(exp, cfg):
        return [r for r in rows if r["Experiment"] == exp and r["Config"] == cfg and r["Status"] == "completed"]

    ALL_CONFIGS = ["X-0000", "X-1000", "X-0100", "X-1100", "X-0010", "X-1010", "X-0110", "X-1110",
                   "X-0001", "X-1001", "X-0101", "X-1101", "X-0011", "X-1011", "X-0111", "X-1111"]
    partial_fill = PatternFill("solid", fgColor="FFF2CC")
    for exp in EXPERIMENTS:
        for cfg in ALL_CONFIGS:
            g = _agg_group(exp, cfg)
            if not g:
                continue
            eff = sorted({str(r["Eff_Batch"]) for r in g})
            row = [exp, cfg, quadrant(cfg), f"{len(g)}/3", ",".join(eff)]
            for col, _label in AGG_METRICS:
                vals = [v for v in (x[col] for x in g) if v not in ("", None)]
                vals = [float(v) for v in vals]
                if not vals:
                    row += ["", "", ""]
                else:
                    m = round(stats.mean(vals), 3)
                    v = round(stats.variance(vals), 4) if len(vals) > 1 else 0.0
                    sd = round(v ** 0.5, 3)
                    row += [m, v, sd]
            ag.append(row)
            if len(g) < 3:
                for c in range(1, len(hdr) + 1):
                    ag.cell(row=ag.max_row, column=c).fill = partial_fill
    ag.freeze_panes = "A4"
    ag.auto_filter.ref = f"A3:{get_column_letter(len(hdr))}{ag.max_row}"
    for i in range(1, len(hdr) + 1):
        ag.column_dimensions[get_column_letter(i)].width = 13
    ag.column_dimensions["A"].width = 11
    ag.column_dimensions["C"].width = 12

    wb.save(path)


def main():
    rows = collect()
    os.makedirs(os.path.join(REPO_ROOT, "results"), exist_ok=True)
    xlsx = os.path.join(REPO_ROOT, "results", "run_tracker.xlsx")
    csvp = os.path.join(REPO_ROOT, "results", "run_tracker.csv")
    write_csv(rows, csvp)
    try:
        write_xlsx(rows, xlsx)
    except Exception as e:
        print(f"[track_runs] xlsx failed ({e}); csv written to {csvp}")
        return
    done = sum(1 for r in rows if r["Status"] == "completed")
    run = sum(1 for r in rows if r["Status"] == "running")
    print(f"[track_runs] {done} completed / {run} running / {len(rows)} total "
          f"-> {xlsx}")


if __name__ == "__main__":
    main()
