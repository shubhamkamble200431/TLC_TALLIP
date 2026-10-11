#!/usr/bin/env python
"""
Human-evaluation COLLECTION template (adequacy/fluency scores left blank for
annotators to fill in) -- NOT a results file. No scores are invented here.

For every (model/config/seed) cell found under prediction_news_data/, draws
the SAME stratified sample of 1000 news_test sentences (200 per length
bucket B1..B5, fixed indices shared across every model so annotators rate
the same inputs each time) and lays out source/prediction/reference with
blank adequacy/fluency columns next to them.

Output: prediction_news_data/human_eval_template.xlsx
    sheet "README"      -- what this is, the 1-5 scale definitions, how to fill it in
    sheet "annotations"  -- the ~208,000-row collection grid
    sheet "models_index"  -- one row per (model,config,seed) cell with its
                             own news_test BLEU, for later consistency checks
                             (NOT used to set any score here)
"""
import json
import os
import random

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = "/raid/kahuja/TLC_TALLIP"
PRED_ROOT = os.path.join(REPO, "prediction_news_data")
RESULTS_ROOT = os.path.join(REPO, "results")  # source/references sidecars live here, not in
                                               # prediction_news_data/ (only "predictions" was copied there)
MANIFEST = os.path.join(PRED_ROOT, "real_test_data", "test_manifest.tsv")
OUT_XLSX = os.path.join(PRED_ROOT, "human_eval_template.xlsx")
N_PER_BUCKET = 200
RNG_SEED = 20261008  # fixed, documented -- same sample for every model


def load_bucket_pool():
    buckets = {}
    with open(MANIFEST) as f:
        header = f.readline().rstrip("\n").split("\t")
        idx_i, bucket_i = header.index("index"), header.index("bucket")
        for line in f:
            parts = line.rstrip("\n").split("\t")
            idx0 = int(parts[idx_i]) - 1
            buckets.setdefault(parts[bucket_i], []).append(idx0)
    return buckets


def sample_indices(buckets):
    rng = random.Random(RNG_SEED)
    sample = {}  # idx0 -> bucket label
    for b in sorted(buckets):
        chosen = rng.sample(buckets[b], min(N_PER_BUCKET, len(buckets[b])))
        for i in chosen:
            sample[i] = b
    return sample  # dict idx0 -> bucket, ~1000 entries


def read_lines(path):
    with open(path, encoding="utf-8") as f:
        return f.read().splitlines()


def find_cells():
    """Every directory under PRED_ROOT/news_test* containing a 'predictions' file."""
    cells = []
    for dirpath, _dirnames, filenames in os.walk(PRED_ROOT):
        if "predictions" not in filenames:
            continue
        if "real_test_data" in dirpath:
            continue
        rel = os.path.relpath(dirpath, PRED_ROOT)
        parts = rel.split(os.sep)
        group = parts[0]                  # news_test | news_test_indic
        seed_name = parts[-1]              # seed_42
        config = parts[-2]                 # X-0000
        model_path = "/".join(parts[1:-2])  # everything else: nllb | tlc_ablation/nllb/huber | nllb/v1 | ...
        cells.append({
            "group": group, "model_path": model_path, "config": config,
            "seed": seed_name.replace("seed_", ""), "dir": dirpath,
            "results_dir": os.path.join(RESULTS_ROOT, rel),  # source/references live here
        })
    return cells


def main():
    buckets = load_bucket_pool()
    sample = sample_indices(buckets)
    sample_idx_sorted = sorted(sample)
    print(f"sampled {len(sample_idx_sorted)} sentence indices "
          f"({ {b: list(sample.values()).count(b) for b in sorted(buckets)} })")

    cells = find_cells()
    print(f"found {len(cells)} (model,config,seed) cells under {PRED_ROOT}")

    rows = []
    models_index_rows = []
    for cell in cells:
        d, rd = cell["dir"], cell["results_dir"]
        preds = read_lines(os.path.join(d, "predictions"))
        srcs = read_lines(os.path.join(rd, "source"))
        refs = read_lines(os.path.join(rd, "references"))
        n = min(len(preds), len(srcs), len(refs))

        mpath = os.path.join(rd, "metrics.json")
        bleu = None
        if os.path.isfile(mpath):
            try:
                bleu = json.load(open(mpath))["metrics"].get("bleu")
            except Exception:
                pass
        models_index_rows.append({
            "group": cell["group"], "model_path": cell["model_path"], "config": cell["config"],
            "seed": cell["seed"], "n_sentences_available": n,
            "news_test_bleu_for_reference_only": round(bleu, 3) if bleu is not None else "",
        })

        for idx0 in sample_idx_sorted:
            if idx0 >= n:
                continue
            rows.append({
                "group": cell["group"], "model_path": cell["model_path"],
                "config": cell["config"], "seed": cell["seed"],
                "bucket": sample[idx0], "sentence_index": idx0 + 1,  # 1-indexed, matches manifest
                "source_sa": srcs[idx0], "prediction": preds[idx0], "reference": refs[idx0],
                "adequacy_1to5": None, "fluency_1to5": None,
            })

    print(f"built {len(rows)} annotation rows")

    df = pd.DataFrame(rows, columns=[
        "group", "model_path", "config", "seed", "bucket", "sentence_index",
        "source_sa", "prediction", "reference", "adequacy_1to5", "fluency_1to5",
    ])
    idx_df = pd.DataFrame(models_index_rows)

    readme_lines = [
        ["Human evaluation COLLECTION template -- adequacy & fluency"],
        [""],
        ["This file has NO scores in it. 'adequacy_1to5' and 'fluency_1to5' are blank"],
        ["for real annotators to fill in. Nothing here was computed or invented."],
        [""],
        ["Sampling: the SAME 1000 news_test sentences (200 per length bucket B1..B5,"],
        [f"fixed RNG seed {RNG_SEED}) are used for every model, so every model is rated"],
        ["on the identical set of source sentences -- only the 'prediction' column differs."],
        [""],
        ["Scales (standard 5-point MT human eval, e.g. WMT-style):"],
        ["  Adequacy (1-5): does the prediction preserve the MEANING of the source/reference?"],
        ["    5 = all meaning preserved   3 = much of the meaning preserved   1 = little/no meaning preserved"],
        ["  Fluency (1-5): is the prediction GRAMMATICAL, natural Hindi (or Gujarati/Marathi)?"],
        ["    5 = flawless language   3 = non-native but comprehensible   1 = incomprehensible"],
        [""],
        ["Rate adequacy and fluency independently -- a fluent sentence can still be inadequate"],
        ["(wrong meaning stated fluently), and an adequate one can be disfluent."],
        [""],
        ["Columns: group/model_path/config/seed identify which checkpoint produced 'prediction';"],
        ["bucket/sentence_index identify which news_test sentence (same across all models);"],
        ["source_sa is the Sanskrit input, reference is the gold translation."],
        [""],
        ["'models_index' sheet lists every (model,config,seed) cell with its own automatic"],
        ["news_test BLEU, included ONLY for your own later cross-checking -- it is NOT used"],
        ["anywhere in this file to set or suggest a score."],
    ]

    ensure_dir = os.path.dirname(OUT_XLSX)
    os.makedirs(ensure_dir, exist_ok=True)

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(readme_lines).to_excel(writer, sheet_name="README", header=False, index=False)
        df.to_excel(writer, sheet_name="annotations", index=False)
        idx_df.to_excel(writer, sheet_name="models_index", index=False)

    # light formatting: header bold, freeze header row, widen key columns
    wb = openpyxl.load_workbook(OUT_XLSX)
    ws = wb["annotations"]
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    widths = {"A": 14, "B": 22, "C": 10, "D": 8, "E": 10, "F": 14,
              "G": 50, "H": 50, "I": 50, "J": 14, "K": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    highlight = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    for col_letter in ("J", "K"):
        for row in range(2, ws.max_row + 1):
            ws[f"{col_letter}{row}"].fill = highlight

    ws2 = wb["models_index"]
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    ws2.freeze_panes = "A2"
    for i, col in enumerate(idx_df.columns, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 2)

    ws0 = wb["README"]
    ws0.column_dimensions["A"].width = 100
    for row in ws0.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.row == 1:
                cell.font = Font(bold=True, size=14)

    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}")
    print(f"  annotations: {len(df)} rows, {len(cells)} models x up to {len(sample_idx_sorted)} sentences")
    print(f"  models_index: {len(idx_df)} rows")


if __name__ == "__main__":
    main()
